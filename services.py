import os
import json
import re
from datetime import timedelta

import jdatetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import settings
from database import get_db
# PERSIAN_MONTHS، PERSIAN_DIGITS و to_persian_digits از models صادر می‌شوند
from models import (
    PERSIAN_DIGITS, PERSIAN_MONTHS,
    Setting, Food, User, Reservation, DinnerReservation, FreeOrder,
    UserContainer, WEEKDAYS_MAP, to_persian_digits,
)


def to_int(text: str) -> int:
    value = str(text).strip()
    if not value:
        raise ValueError("empty numeric value")
    value = value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))
    value = value.replace("،", "").replace(" ", "")
    if value.startswith("+"):
        value = value[1:]
    if value.startswith("-"):
        raise ValueError("negative values are not allowed")
    if not re.fullmatch(r"\d+", value):
        raise ValueError("invalid numeric value")
    return int(value)


def format_price(amount: int) -> str:
    return f"{amount:,}"


def get_setting(key: str) -> str:
    with get_db() as db:
        row = db.query(Setting).filter(Setting.key == key).first()
        return row.value if row else ""


def set_setting(key: str, value: str):
    with get_db() as db:
        row = db.query(Setting).filter(Setting.key == key).first()
        if row:
            row.value = str(value)
        else:
            db.add(Setting(key=key, value=str(value)))


def get_all_admin_ids() -> list[str]:
    """آیدی ادمین اصلی (از .env) به‌همراه ادمین‌های اضافه‌شده توسط ادمین اصلی را برمی‌گرداند."""
    ids = list(settings.admin_ids_list)
    extra = get_setting("extra_admin_ids")
    for item in extra.split(","):
        item = item.strip()
        if item and item not in ids:
            ids.append(item)
    return ids


def add_extra_admin(admin_id: str) -> bool:
    """ادمین جدید اضافه می‌کند؛ اگر قبلاً بوده False برمی‌گرداند."""
    admin_id = admin_id.strip()
    if not admin_id.isdigit():
        return False
    if admin_id in get_all_admin_ids():
        return False
    current = [item for item in get_setting("extra_admin_ids").split(",") if item.strip()]
    current.append(admin_id)
    set_setting("extra_admin_ids", ",".join(current))
    return True


def remove_extra_admin(admin_id: str) -> bool:
    """ادمین اضافه را حذف می‌کند؛ ادمین اصلی (.env) قابل حذف نیست."""
    admin_id = admin_id.strip()
    current = [item for item in get_setting("extra_admin_ids").split(",") if item.strip()]
    if admin_id not in current:
        return False
    current.remove(admin_id)
    set_setting("extra_admin_ids", ",".join(current))
    return True


def get_registration_period() -> tuple[int, int]:
    now = jdatetime.datetime.now()
    year = get_setting("registration_year")
    month = get_setting("registration_month")
    if year.isdigit() and month.isdigit():
        return int(year), int(month)
    return now.year, now.month


def get_registration_window() -> tuple[int, int]:
    try:
        start_day = int(get_setting("registration_start_day") or 1)
        end_day = int(get_setting("registration_end_day") or 7)
    except ValueError:
        return 1, 7
    if not 1 <= start_day <= end_day <= 31:
        return 1, 7
    return start_day, end_day


def get_registration_window_config() -> tuple[int, int, int, int]:
    start_day, end_day = get_registration_window()
    try:
        start_offset = int(get_setting("registration_start_month_offset") or 0)
        end_offset = int(get_setting("registration_end_month_offset") or 0)
    except ValueError:
        return start_day, end_day, 0, 0
    if start_offset > end_offset or start_offset < -12 or end_offset > 12:
        return start_day, end_day, 0, 0
    return start_day, end_day, start_offset, end_offset


def shift_period(year: int, month: int, offset: int) -> tuple[int, int]:
    absolute_month = year * 12 + month - 1 + offset
    return absolute_month // 12, absolute_month % 12 + 1


def is_registration_window_day() -> bool:
    start_day, end_day, start_offset, end_offset = get_registration_window_config()
    target_year, target_month = get_registration_period()
    start_year, start_month = shift_period(target_year, target_month, start_offset)
    end_year, end_month = shift_period(target_year, target_month, end_offset)
    today = jdatetime.datetime.now()
    today_key = (today.year, today.month, today.day)
    return (start_year, start_month, start_day) <= today_key <= (end_year, end_month, end_day)


def is_registration_open() -> bool:
    if get_setting("registration_status") != "active":
        return False
    if get_setting("registration_manual_override") == "true":
        return True
    return is_registration_window_day()


def registration_deadline_text() -> str:
    _, end_day, _, end_offset = get_registration_window_config()
    target_year, target_month = get_registration_period()
    end_year, end_month = shift_period(target_year, target_month, end_offset)
    return f"روز {end_day} ماه {PERSIAN_MONTHS[end_month - 1]}"


def registration_period_text() -> str:
    _, month = get_registration_period()
    return PERSIAN_MONTHS[month - 1]


def registration_days_left() -> int:
    """روزهای باقی‌مانده تا پایان مهلت ثبت‌نام (حداقل صفر)."""
    _, end_day, _, end_offset = get_registration_window_config()
    target_year, target_month = get_registration_period()
    end_year, end_month = shift_period(target_year, target_month, end_offset)
    if end_month <= 6:
        days_in_month = 31
    elif end_month <= 11:
        days_in_month = 30
    else:
        days_in_month = 30 if jdatetime.date(end_year, end_month, 1).is_leap_year() else 29
    end_date = jdatetime.date(end_year, end_month, min(end_day, days_in_month))
    return max(0, (end_date - jdatetime.date.today()).days)


def registration_deadline_notice() -> str:
    """یک‌خط خلاصه مهلت برای نمایش زیر منوی اصلی."""
    days_left = registration_days_left()
    if days_left == 0:
        remaining = "امروز آخرین روز است"
    else:
        remaining = f"{to_persian_digits(days_left)} روز باقی‌مانده"
    return f"⏳ مهلت ثبت‌نام: {registration_deadline_text()} — {remaining}"


def registration_window_text() -> str:
    """بازه ثبت‌نام را به‌صورت خوانا با نام ماه‌ها برمی‌گرداند."""
    start_day, end_day, start_offset, end_offset = get_registration_window_config()
    target_year, target_month = get_registration_period()
    start_year, start_month = shift_period(target_year, target_month, start_offset)
    end_year, end_month = shift_period(target_year, target_month, end_offset)
    if (start_year, start_month) == (end_year, end_month):
        return f"از روز {start_day} تا روز {end_day} {PERSIAN_MONTHS[end_month - 1]}"
    return (
        f"از روز {start_day} {PERSIAN_MONTHS[start_month - 1]} "
        f"تا روز {end_day} {PERSIAN_MONTHS[end_month - 1]}"
    )


def format_jalali_datetime(value) -> str | None:
    """datetime ذخیره‌شده در دیتابیس را به رشته‌ی شمسی خوانا تبدیل می‌کند."""
    if not value:
        return None
    try:
        local_dt = value + timedelta(hours=3, minutes=30)
        jd = jdatetime.datetime.fromgregorian(datetime=local_dt)
        return f"{jd.year}/{jd.month:02d}/{jd.day:02d} - {jd.hour:02d}:{jd.minute:02d}"
    except (TypeError, ValueError, OverflowError):
        return None


def admin_notice_text(title: str, actor_id: str | None = None, note: str = "") -> str:
    """اعلان ادمین‌ها را با قالب ثابت (زمان، دوره، انجام‌دهنده) می‌سازد."""
    now = jdatetime.datetime.now()
    stamp = f"{now.year}/{now.month:02d}/{now.day:02d} - {now.hour:02d}:{now.minute:02d}"
    lines = [title, f"🕒 زمان: {stamp} | 📅 دوره: {registration_period_text()}"]
    if actor_id:
        lines.append(f"👤 انجام‌دهنده: {actor_id}")
    if note:
        lines.append(note)
    return "\n".join(lines)


def get_next_period(year: int, month: int) -> tuple[int, int]:
    return (year + 1, 1) if month == 12 else (year, month + 1)


def get_days_details() -> list:
    year, month = get_registration_period()

    if month <= 6:
        days_in_month = 31
    elif month <= 11:
        days_in_month = 30
    else:
        days_in_month = 30 if jdatetime.date(year, 12, 1).is_leap_year() else 29

    holidays_str = get_setting("holidays")
    holidays = [int(x.strip()) for x in holidays_str.split(",") if x.strip().isdigit()]

    blocked_str = get_setting("blocked_days")
    blocked = [int(x.strip()) for x in blocked_str.split(",") if x.strip().isdigit()]

    details = []
    for d in range(1, days_in_month + 1):
        jd = jdatetime.date(year, month, d)
        w_day = jd.weekday()
        details.append({
            "day": d,
            "weekday_num": w_day,
            "weekday_name": WEEKDAYS_MAP[w_day]["name"],
            "is_holiday": (w_day == 6) or (d in holidays) or (d in blocked),
        })
    return details


def calculate_user_total(user_id: str, meal_type: str | None = None) -> int:
    with get_db() as db:
        days_details = {d["day"]: d for d in get_days_details()}

        meal_types = [meal_type] if meal_type else ["lunch", "dinner"]
        normal_total = 0
        free_total = 0
        for current_meal in meal_types:
            reservation_model = Reservation if current_meal == "lunch" else DinnerReservation
            res_days = db.query(reservation_model).filter(
                reservation_model.user_id == user_id
            ).all()
            for reservation in res_days:
                day = days_details.get(reservation.day)
                if day and not day["is_holiday"]:
                    food = db.query(Food).filter(
                        Food.day_num == day["weekday_num"],
                        Food.meal_type == current_meal,
                    ).first()
                    if food:
                        normal_total += food.normal_price

        free_orders = db.query(FreeOrder, Food).join(
            Food, FreeOrder.food_id == Food.id
        ).filter(
            FreeOrder.user_id == user_id,
            Food.meal_type.in_(meal_types),
        ).all()
        free_total = sum(
            food.free_price * order.count
            for order, food in free_orders
        )

        container_total = 0
        if meal_type in (None, "lunch"):
            container = db.query(UserContainer).filter(
                UserContainer.user_id == user_id
            ).first()
            container_total = (container.count * int(get_setting("container_price"))) if container else 0

    return normal_total + free_total + container_total


def get_report_snapshot() -> dict:
    year, month = get_registration_period()
    days = get_days_details()
    with get_db() as db:
        foods = [
            {"id": f.id, "day_num": f.day_num, "day_name": f.day_name,
             "food_name": f.food_name, "normal_price": f.normal_price,
             "free_price": f.free_price, "meal_type": f.meal_type}
            for f in db.query(Food).order_by(Food.id).all()
        ]
        users = []
        for user in db.query(User).all():
            users.append({
                "user_id": user.user_id,
                "full_name": user.full_name,
                "status": user.status,
                "lunch_status": user.lunch_status,
                "dinner_status": user.dinner_status,
                "reservations": [r.day for r in db.query(Reservation).filter(Reservation.user_id == user.user_id).all()],
                "dinner_reservations": [r.day for r in db.query(DinnerReservation).filter(DinnerReservation.user_id == user.user_id).all()],
                "free_orders": [
                    {"food_id": order.food_id, "count": order.count}
                    for order in db.query(FreeOrder).filter(FreeOrder.user_id == user.user_id).all()
                ],
                "containers": next((c.count for c in db.query(UserContainer).filter(UserContainer.user_id == user.user_id).all()), 0),
            })
        container_price = int((db.query(Setting).filter(Setting.key == "container_price").first() or Setting(value="0")).value or 0)
        return {"year": year, "month": month, "days": days, "foods": foods,
            "users": users, "container_price": container_price}


def _snapshot_user_total(user: dict, snapshot: dict) -> int:
    days = {day["day"]: day for day in snapshot["days"]}
    foods = {food["id"]: food for food in snapshot["foods"]}
    total = 0
    for meal_type, reservation_key in (("lunch", "reservations"), ("dinner", "dinner_reservations")):
        for day_number in user.get(reservation_key, []):
            day = days.get(day_number)
            if day and not day["is_holiday"]:
                food = next((item for item in snapshot["foods"] if item["day_num"] == day["weekday_num"] and item["meal_type"] == meal_type), None)
                total += food["normal_price"] if food else 0
    for order in user["free_orders"]:
        food = foods.get(order["food_id"])
        total += (food["free_price"] * order["count"]) if food else 0
    return total + user["containers"] * snapshot["container_price"]


def build_excel_report(snapshot: dict | None = None, file_suffix: str = "") -> str:
    snapshot = snapshot or get_report_snapshot()
    wb = openpyxl.Workbook()
    font_family = "Tahoma"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font   = Font(name=font_family, size=11)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    alt_fill    = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),  bottom=Side(style="thin", color="D9D9D9"),
    )

    year = snapshot["year"]
    month = snapshot["month"]
    month_name = PERSIAN_MONTHS[month - 1]
    period = f"{month_name} {to_persian_digits(year)}"
    days = snapshot["days"]

    # شیت آشپزخانه
    ws1 = wb.active
    ws1.title = f"آشپزخانه {period}"[:31]
    ws1.sheet_view.rightToLeft = True
    ws1.append([f"موسسه امام حسین (ع) - گزارش آشپزخانه {period}"])
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    ws1.append(["وعده", "روز ماه", "روز هفته", "نام غذا", "تعداد ثبت‌شده"])

    for meal_type, meal_name in (("lunch", "ناهار"), ("dinner", "شام")):
        meal_foods = [food for food in snapshot["foods"] if food["meal_type"] == meal_type]
        reservation_key = "reservations" if meal_type == "lunch" else "dinner_reservations"
        for day in days:
            food = next((item for item in meal_foods if item["day_num"] == day["weekday_num"]), None)
            count = sum(day["day"] in user[reservation_key] for user in snapshot["users"])
            ws1.append([meal_name, day["day"], day["weekday_name"], "تعطیل رسمی / جمعه" if day["is_holiday"] else (food["food_name"] if food else "نامشخص"), 0 if day["is_holiday"] else count])

    # شیت مالی
    ws2 = wb.create_sheet(title=f"مالی {period}"[:31])
    ws2.sheet_view.rightToLeft = True

    ws2.append([f"موسسه امام حسین (ع) - گزارش مالی {period}"])
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(snapshot["foods"]) * 2 + 1)
    headers = ["نام و نام خانوادگی", "قیمت کل (تومان)"]
    for food in snapshot["foods"]:
        meal_name = "ناهار" if food["meal_type"] == "lunch" else "شام"
        headers += [f"{meal_name} عادی {food['food_name']}", f"{meal_name} آزاد {food['food_name']}"]
    headers.append("ظرف یکبار مصرف")
    ws2.append(headers)
    days_details = {day["day"]: day for day in days}
    for user in snapshot["users"]:
        if user["status"] != "completed":
            continue
        food_counts = {food["id"]: {"norm": 0, "free": 0} for food in snapshot["foods"]}
        for meal_type, reservation_key in (("lunch", "reservations"), ("dinner", "dinner_reservations")):
            for day_number in user[reservation_key]:
                day = days_details.get(day_number)
                if day and not day["is_holiday"]:
                    food = next((item for item in snapshot["foods"] if item["day_num"] == day["weekday_num"] and item["meal_type"] == meal_type), None)
                    if food:
                        food_counts[food["id"]]["norm"] += 1
        for order in user["free_orders"]:
            if order["food_id"] in food_counts:
                food_counts[order["food_id"]]["free"] = order["count"]
        row = [user["full_name"], _snapshot_user_total(user, snapshot)]
        for food in snapshot["foods"]:
            row += [food_counts[food["id"]]["norm"], food_counts[food["id"]]["free"]]
        row.append(user["containers"])
        ws2.append(row)

    # استایل
    for ws in [ws1, ws2]:
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column)):
            for cell in row:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                elif row_idx == 0 and cell.value:
                    cell.font = Font(name=font_family, size=13, bold=True, color="1F4E78")
                    cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
                elif row_idx > 1 and row_idx % 2 == 0:
                    cell.fill = alt_fill
                if "مالی" in ws.title and row_idx > 1 and cell.column == 2:
                    cell.number_format = "#,##0"
        for column_cells in ws.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 3, 12), 35)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = ws.dimensions

    file_name = f"Food_Report_{year}_{month:02d}{file_suffix}.xlsx"
    wb.save(file_name)
    return file_name


def get_menu_text() -> str:
    with get_db() as db:
        # جمعه ناهار ندارد و پنج‌شنبه شام ندارد؛ از منو حذف می‌شوند
        foods = [
            (f.meal_type, f.day_name, f.food_name, f.normal_price, f.free_price)
            for f in db.query(Food).filter(
                ~((Food.meal_type == "lunch") & (Food.day_num == 6)),
                ~((Food.meal_type == "dinner") & (Food.day_num == 5)),
            ).order_by(Food.day_num).all()
        ]
    c_price = get_setting("container_price")
    lines = [f"{registration_deadline_notice()}\n"]
    lines.append("📊 لیست منوی هفتگی و قیمت‌ها (تومان):\n")
    lunch_foods = [food for food in foods if food[0] == "lunch"]
    dinner_foods = [food for food in foods if food[0] == "dinner"]
    lines.append("🍽️ ناهار")
    for _, day_name, food_name, normal_price, free_price in lunch_foods:
        lines.append(f"🔹 {day_name}: {food_name}")
        lines.append(f"   💵 قیمت عادی: {format_price(normal_price)} | آزاد: {format_price(free_price)}")
    lines.append(f"\n📦 قیمت ظرف یکبار مصرف ناهار: {format_price(int(c_price))} تومان")
    lines.append("\n🌙 شام")
    for _, day_name, food_name, normal_price, free_price in dinner_foods:
        lines.append(f"🔹 {day_name}: {food_name}")
        lines.append(f"   💵 قیمت عادی: {format_price(normal_price)} | آزاد: {format_price(free_price)}")
    lines.append("--------------------------------------")
    lines.append("میتوانید اقدام به ثبت یا ویرایش کنید")
    lines.append(registration_deadline_notice())
    return "\n".join(lines)


def build_invoice_text(user_id: str, meal_type: str = "lunch") -> str:
    total = calculate_user_total(user_id, meal_type)
    with get_db() as db:
        reservation_model = Reservation if meal_type == "lunch" else DinnerReservation
        res_rows = db.query(reservation_model).filter(reservation_model.user_id == user_id).all()
        res_days = sorted([r.day for r in res_rows])
        first_created = min((r.created_at for r in res_rows if r.created_at), default=None)
        last_updated = max(
            ((r.updated_at or r.created_at) for r in res_rows if (r.updated_at or r.created_at)),
            default=None,
        )
        days_details = {item["day"]: item for item in get_days_details()}
        normal_days = []
        normal_total = 0
        for day_number in res_days:
            day_info = days_details.get(day_number)
            food = db.query(Food).filter(
                Food.meal_type == meal_type,
                Food.day_num == (day_info["weekday_num"] if day_info else -1),
            ).first() if day_info else None
            if day_info:
                price = food.normal_price if food else 0
                normal_days.append(to_persian_digits(day_number))
                normal_total += price
        free_orders = [
            (fo.count, food.food_name, food.free_price)
            for fo, food in (
                db.query(FreeOrder, Food)
                .join(Food, FreeOrder.food_id == Food.id)
                .filter(FreeOrder.user_id == user_id, Food.meal_type == meal_type)
                .all()
            )
        ]
        c = db.query(UserContainer).filter(UserContainer.user_id == user_id).first() if meal_type == "lunch" else None
        cnt = c.count if c else 0
        free_total = sum(count * price for count, _, price in free_orders)
        container_total = cnt * int(get_setting("container_price")) if meal_type == "lunch" else 0

    meal_name = "ناهار" if meal_type == "lunch" else "شام"
    has_registration = bool(normal_days or any(count > 0 for count, _, _ in free_orders) or (meal_type == "lunch" and cnt > 0))
    status_line = (
        f"✅ {meal_name} شما با موفقیت ثبت شده است.\n"
        if has_registration
        else f"ℹ️ {meal_name}ی برای شما ثبت نشده است.\n"
    )
    lines = [
        status_line,
        f"📅 تاریخ روزهای عادی: {', '.join(normal_days) if normal_days else 'هیچ روزی ثبت نشده'}",
        f"💰 کل روزهای عادی: {format_price(normal_total)} تومان",
    ]
    lines.append(f"\n🍔 غذاهای آزاد {'ناهار' if meal_type == 'lunch' else 'شام'}:")
    for count, food_name, free_price in free_orders:
        if count > 0:
            lines.append(f"   - {food_name}: {count} عدد")
    lines.append(f"💰 کل غذاهای آزاد: {format_price(free_total)} تومان")
    if meal_type == "lunch":
        lines.append(f"\n📦 تعداد ظروف یکبار مصرف: {cnt} عدد")
        lines.append(f"💰 هزینه ظروف: {format_price(container_total)} تومان")
    created_text = format_jalali_datetime(first_created)
    updated_text = format_jalali_datetime(last_updated)
    if created_text:
        lines.append(f"🕒 تاریخ ثبت: {created_text}")
    if updated_text and updated_text != created_text:
        lines.append(f"✏️ آخرین ویرایش: {updated_text}")
    lines += ["--------------------------------------", f"💰 قیمت کل {'ناهار' if meal_type == 'lunch' else 'شام'}: {format_price(total)} تومان"]
    return "\n".join(lines)
